import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def build_docx(data: dict, lang: str) -> bytes:
    doc = Document()

    title = data.get("title") or "Report"
    intro = data.get("introduction") or ""
    sections = data.get("sections") or []
    conclusion = data.get("conclusion") or ""
    references = data.get("references") or []

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(22)

    doc.add_heading("Introduction" if lang == "en" else "المقدمة", level=1)
    doc.add_paragraph(intro)

    for sec in sections:
        heading = sec.get("heading") or ""
        content = sec.get("content") or ""
        doc.add_heading(heading, level=1)
        doc.add_paragraph(content)

    doc.add_heading("Conclusion" if lang == "en" else "الخاتمة", level=1)
    doc.add_paragraph(conclusion)

    doc.add_heading("References" if lang == "en" else "المصادر", level=1)
    if references:
        for ref in references:
            doc.add_paragraph(ref, style="List Number")
    else:
        doc.add_paragraph("-")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def preview_text(data: dict) -> str:
    title = data.get("title") or ""
    intro = data.get("introduction") or ""
    sections = data.get("sections") or []
    headings = "\n".join(f"• {s.get('heading', '')}" for s in sections)
    intro_short = intro[:500] + ("…" if len(intro) > 500 else "")
    text = f"📄 {title}\n\n{intro_short}"
    if headings:
        text += f"\n\n📚 {headings}"
    return text